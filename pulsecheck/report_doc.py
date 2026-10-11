"""Build the technical report (Word .docx) from the code and the real result files.

    python -m pulsecheck report-doc --author "Your Name"

Every number in the Results chapter is read from reports/*.json, which only the
real-dataset commands write. A section whose report does not exist yet says so
and gives the command that produces it - the document never invents a result.
Library versions are read from the installed packages, so the stack table is
always accurate. Re-run the command after each training run.
"""

from __future__ import annotations

import importlib.metadata as md
import json
import platform
import re
import tempfile
import time
from pathlib import Path

import numpy as np

from . import __version__
from .config import (DETECT_STEP_S, DETECT_WINDOW_S, HR_BAND, MODELS_DIR, REPORTS_DIR, ROOT, RPPG_WINDOW_S,
                     TARGET_FS, WIDE_BAND)
from .features import describe, feature_names
from .roi import GRID_LANDMARKS, ROI_LANDMARKS, ROI_NAMES

# Chart colours: the website's validated palette (light mode), so figures match the site.
S = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
INK, INK2, MUTED, GRID, AXIS, DEEMPH = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7", "#b4b2aa"
METHOD = {"pos": "POS", "chrom": "CHROM", "green": "Green"}
MANIP = ["Deepfakes", "Face2Face", "FaceSwap", "NeuralTextures", "FaceShifter"]


def _load(name):
    p = REPORTS_DIR / name
    if not p.exists():
        return None
    d = json.loads(p.read_text())
    return None if d.get("source") == "synthetic" else d  # belt and braces: never report synthetic runs


def _f(v, d=2):
    return "-" if v is None or (isinstance(v, float) and not np.isfinite(v)) else f"{v:.{d}f}"


def _pct(v, d=1):
    return "-" if v is None else f"{100 * v:.{d}f}%"


def _ver(pkg):
    try:
        return md.version(pkg)
    except md.PackageNotFoundError:
        return None


# --------------------------------------------------------------------------- figures (matplotlib)

class Figs:
    def __init__(self, folder: Path):
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        self.plt = plt
        self.dir = folder
        self.n = 0
        plt.rcParams.update({
            "font.family": "sans-serif", "font.sans-serif": ["Segoe UI", "Calibri", "DejaVu Sans"], "font.size": 9,
            "axes.edgecolor": AXIS, "axes.labelcolor": INK2, "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8,
            "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False,
            "axes.titlesize": 10, "axes.titlecolor": INK, "legend.frameon": False, "lines.linewidth": 2,
            "savefig.dpi": 200, "savefig.bbox": "tight",
        })

    def new(self, w=6.4, h=3.0, **kw):
        return self.plt.subplots(figsize=(w, h), **kw)

    def save(self, fig) -> Path:
        self.n += 1
        p = self.dir / f"fig{self.n:02d}.png"
        fig.savefig(p)
        self.plt.close(fig)
        return p

    # ---- schematic
    def pipeline(self):
        from matplotlib.patches import FancyBboxPatch

        fig, ax = self.new(6.6, 3.96)
        ax.axis("off")
        ax.set_xlim(0, 100)
        ax.set_ylim(0, 60)
        W, H = 21, 9

        def box(cx, cy, text, face, edge, w=W):
            ax.add_patch(FancyBboxPatch((cx - w / 2, cy - H / 2), w, H, boxstyle="round,pad=0,rounding_size=1.2",
                                        facecolor=face, edgecolor=edge, lw=1.1))
            ax.text(cx, cy, text, ha="center", va="center", fontsize=7.6, color=INK, linespacing=1.3)

        def arrow(x1, y1, x2, y2):
            ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="-|>", color=INK2, lw=0.9, mutation_scale=8))

        blue, orange, green = ("#e6effb", S[0]), ("#fdeee6", S[1]), ("#e3f4ec", S[2])
        row1 = [(12, "Video file"), (37, "Face landmarks\nMediaPipe, 478 points"), (62, "Skin patches\n6 regions + 32 grid"),
                (87, "Colour traces\nRGB per patch, 30 Hz")]
        for (x, t) in row1:
            box(x, 50, t, *blue)
        for (a, _), (b, _) in zip(row1[:-1], row1[1:]):
            arrow(a + W / 2, 50, b - W / 2, 50)
        box(20, 33, "Pulse (POS / CHROM / Green)\nfilter + FFT -> heart rate", *blue, w=30)
        box(55, 33, "91 agreement features\nper 8 s window", *orange)
        box(82, 33, "PPG maps\n2 x 32 x 224", *orange)
        for x in (20, 55, 82):
            arrow(87, 45.5, x, 37.5)
        box(55, 19, "Classical model\nscikit-learn (CPU)", *orange)
        box(82, 19, "CNN\nPyTorch (GPU)", *orange)
        arrow(55, 28.5, 55, 23.5)
        arrow(82, 28.5, 82, 23.5)
        box(68.5, 5.5, "Window votes -> video score -> verdict", *green, w=40)
        arrow(55, 14.5, 62, 10)
        arrow(82, 14.5, 75, 10)
        return self.save(fig)

    # ---- heart rate
    def hr_scatter_ba(self, hr):
        sc, hl = hr["scatter"], hr["headline"]
        gt, est = np.array(sc["gt"]), np.array(sc["est"])
        fig, (a, b) = self.new(6.6, 2.9, ncols=2)
        lo, hi = min(gt.min(), est.min()) - 5, max(gt.max(), est.max()) + 5
        a.plot([lo, hi], [lo, hi], color=AXIS, lw=1)
        a.scatter(gt, est, s=16, color=S[0], edgecolor="white", linewidth=0.6, zorder=3)
        a.set(xlim=(lo, hi), ylim=(lo, hi), xlabel="oximeter, bpm", ylabel="video, bpm", title="Estimate vs oximeter")
        m, d = (gt + est) / 2, est - gt
        b.scatter(m, d, s=16, color=S[0], edgecolor="white", linewidth=0.6, zorder=3)
        for y, c in ((hl["bias"], INK2), (hl["loa_low"], AXIS), (hl["loa_high"], AXIS)):
            b.axhline(y, color=c, lw=1.2 if c == INK2 else 1)
        b.set(xlabel="mean of the two, bpm", ylabel="video - oximeter, bpm", title="Bland-Altman")
        fig.tight_layout()
        return self.save(fig)

    def hr_windows(self, hr):
        fig, ax = self.new(4.8, 2.7)
        for i, m in enumerate(hr["methods"]):
            rows = sorted((s for s in hr["summary"] if s["method"] == m), key=lambda s: s["window_s"])
            ax.plot([r["window_s"] for r in rows], [r["mae"] for r in rows], "-o", color=S[i], ms=5, mec="white", label=METHOD[m])
        ax.set(xlabel="window length, seconds", ylabel="MAE, bpm", title="Error vs clip length", xticks=hr["windows_s"])
        ax.set_ylim(bottom=0)
        ax.legend()
        return self.save(fig)

    def hr_subjects(self, hr):
        ps = hr["per_subject"]
        fig, ax = self.new(6.4, 2.5)
        ax.bar(range(len(ps)), [p["mae"] for p in ps], color=S[0], width=0.7)
        ax.set(xlabel="subjects, sorted", ylabel="MAE, bpm", title="Error per person", xticks=[])
        return self.save(fig)

    def compression(self, comp):
        rows = [r for r in comp["rows"] if r["crf"] is not None]
        base = next((r for r in comp["rows"] if r["crf"] is None), None)
        fig, ax = self.new(4.8, 2.7)
        x = [r["crf"] for r in rows]
        ax.plot(x, [r["mae_pos_10s"] for r in rows], "-o", color=S[0], ms=5, mec="white", label="POS")
        ax.plot(x, [r["mae_green_10s"] for r in rows], "-o", color=S[2], ms=5, mec="white", label="Green")
        if base and base["mae_pos_10s"] is not None:
            ax.axhline(base["mae_pos_10s"], color=AXIS, lw=1)
            ax.text(x[0], base["mae_pos_10s"], " uncompressed (POS)", va="bottom", fontsize=7, color=INK2)
        ax.set(xlabel="H.264 CRF (higher = more compressed)", ylabel="MAE, bpm (10 s)", title="Compression vs error", xticks=x)
        ax.set_ylim(bottom=0)
        ax.legend()
        return self.save(fig)

    # ---- detector
    def rocs(self, det, cnn):
        fig, (a, b) = self.new(6.6, 3.0, ncols=2)
        for ax in (a, b):
            ax.plot([0, 1], [0, 1], color=AXIS, lw=1)
            ax.set(xlim=(0, 1), ylim=(0, 1.02), xlabel="false-alarm rate", ylabel="detection rate")
        if det:
            r = det["test_video"]["roc"]
            a.step(r["x"], r["y"], where="post", color=S[0], label=f"{det['chosen_label']} (AUC {_f(det['test_video']['auc'], 3)})")
        if cnn:
            r = cnn["test_video"]["roc"]
            a.step(r["x"], r["y"], where="post", color=S[1], label=f"CNN (AUC {_f(cnn['test_video']['auc'], 3)})")
        a.set_title("Test ROC, all manipulations")
        a.legend(loc="lower right", fontsize=7)
        src = det or cnn
        for p in sorted(src["per_method"], key=lambda p: MANIP.index(p["method"]) if p["method"] in MANIP else 9):
            i = MANIP.index(p["method"]) if p["method"] in MANIP else 7
            b.step(p["roc"]["x"], p["roc"]["y"], where="post", color=S[i], lw=1.6, label=f"{p['method']} ({_f(p['auc'], 2)})")
        b.set_title(f"By manipulation ({'classical' if det else 'CNN'})")
        b.legend(loc="lower right", fontsize=6.5)
        fig.tight_layout()
        return self.save(fig)

    def importance(self, det):
        imp = det["importance"][:12][::-1]
        fig, ax = self.new(5.6, 3.2)
        ax.barh([d["label"] for d in imp], [d["auc_drop"] for d in imp], xerr=[d["std"] for d in imp], color=S[0],
                height=0.6, error_kw=dict(ecolor=AXIS, lw=1))
        ax.set(xlabel="drop in test AUC when shuffled", title="Permutation importance (top 12)")
        ax.grid(axis="y", visible=False)
        return self.save(fig)

    def learning(self, cnn):
        h = cnn["history"]
        ep = [r["epoch"] for r in h]
        fig, (a, b) = self.new(6.6, 2.6, ncols=2)
        a.plot(ep, [r["train_loss"] for r in h], color=S[0])
        a.set(xlabel="epoch", ylabel="loss", title="Training loss")
        b.plot(ep, [r["val_video_auc"] for r in h], color=S[1])
        b.axvline(cnn["training"]["best_epoch"], color=AXIS, lw=1)
        b.set(xlabel="epoch", ylabel="AUC", title="Validation video AUC (early stopping)")
        fig.tight_layout()
        return self.save(fig)


# --------------------------------------------------------------------------- docx helpers

class Doc:
    def __init__(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        self.Inches, self.Pt, self.RGB, self.qn, self.Ox, self.ALIGN = Inches, Pt, RGBColor, qn, OxmlElement, WD_ALIGN_PARAGRAPH
        self.d = Document()
        st = self.d.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(10.5)
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.12
        for lvl, size, color in ((1, 16, "1C5CAB"), (2, 13, "256ABF"), (3, 11, "3B3A37")):
            h = st[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = RGBColor.from_string(color)
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(4)
        for s in self.d.sections:
            s.left_margin = s.right_margin = Inches(1)
            s.top_margin = s.bottom_margin = Inches(0.9)
        self.fig_n = 0
        self.tab_n = 0

    # text
    def h(self, text, level=1):
        return self.d.add_heading(text, level)

    def p(self, *parts, style=None, italic=False, size=None, color=None, align=None):
        """Paragraph from strings; wrap a string in ** ** for bold, ` ` for code."""
        para = self.d.add_paragraph(style=style)
        for part in parts:
            self._runs(para, part, italic=italic, size=size, color=color)
        if align == "center":
            para.alignment = self.ALIGN.CENTER
        return para

    def bullets(self, items, numbered=False):
        for it in items:
            self.p(it, style="List Number" if numbered else "List Bullet")

    def code(self, text):
        for line in text.strip("\n").splitlines():
            para = self.d.add_paragraph()
            para.paragraph_format.space_after = self.Pt(0)
            para.paragraph_format.left_indent = self.Inches(0.15)
            r = para.add_run(line if line else " ")
            r.font.name = "Consolas"
            r.font.size = self.Pt(8.5)
            self._shade(para, "F3F2EE")
        self.d.add_paragraph().paragraph_format.space_after = self.Pt(2)

    def _shade(self, para, fill):
        pPr = para._p.get_or_add_pPr()
        shd = self.Ox("w:shd")
        shd.set(self.qn("w:val"), "clear")
        shd.set(self.qn("w:color"), "auto")
        shd.set(self.qn("w:fill"), fill)
        pPr.append(shd)

    def _runs(self, para, text, italic=False, size=None, color=None):
        """Add ``text`` to ``para``; **bold** and `code` spans are formatted."""
        for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not chunk:
                continue
            if chunk.startswith("**"):
                r = para.add_run(chunk[2:-2])
                r.bold = True
            elif chunk.startswith("`"):
                r = para.add_run(chunk[1:-1])
                r.font.name = "Consolas"
                r.font.size = self.Pt(9.5)
            else:
                r = para.add_run(chunk)
            r.italic = italic or r.italic
            if size:
                r.font.size = self.Pt(size)
            if color:
                r.font.color.rgb = self.RGB.from_string(color)

    def note(self, title, text, fill="FFF4D6"):
        t = self.d.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        self._cell_fill(cell, fill)
        cp = cell.paragraphs[0]
        r = cp.add_run(title + " ")
        r.bold = True
        self._runs(cp, text)
        self.d.add_paragraph().paragraph_format.space_after = self.Pt(0)

    def pending(self, what, commands):
        self.note("Not run yet.", f"{what} will appear here after you run:", fill="F3F2EE")
        self.code("\n".join(commands))

    # tables
    def _cell_fill(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = self.Ox("w:shd")
        shd.set(self.qn("w:val"), "clear")
        shd.set(self.qn("w:color"), "auto")
        shd.set(self.qn("w:fill"), fill)
        tcPr.append(shd)

    def table(self, header, rows, widths=None, caption=None, font=9):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(htxt))
            r.bold = True
            r.font.size = self.Pt(font)
            self._cell_fill(c, "DCE8F7")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run("-" if v is None else str(v))
                r.font.size = self.Pt(font)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = self.Inches(w)
        if caption:
            self.tab_n += 1
            cp = self.p(f"Table {self.tab_n}. {caption}", style="Caption")
            cp.paragraph_format.space_after = self.Pt(10)
        else:
            self.d.add_paragraph().paragraph_format.space_after = self.Pt(2)
        return t

    def figure(self, path, caption, width=6.3):
        self.d.add_picture(str(path), width=self.Inches(width))
        self.d.paragraphs[-1].alignment = self.ALIGN.CENTER
        self.fig_n += 1
        cp = self.p(f"Figure {self.fig_n}. {caption}", style="Caption")
        cp.alignment = self.ALIGN.CENTER

    def link(self, para, url, text):
        part = self.d.part
        r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        hl = self.Ox("w:hyperlink")
        hl.set(self.qn("r:id"), r_id)
        run = self.Ox("w:r")
        rPr = self.Ox("w:rPr")
        c = self.Ox("w:color")
        c.set(self.qn("w:val"), "1C5CAB")
        u = self.Ox("w:u")
        u.set(self.qn("w:val"), "single")
        rPr.append(c)
        rPr.append(u)
        run.append(rPr)
        t = self.Ox("w:t")
        t.text = text
        run.append(t)
        hl.append(run)
        para._p.append(hl)

    def field(self, para, instr, placeholder=""):
        run = para.add_run()
        for kind, text in (("begin", None), (None, instr), ("separate", None), (None, placeholder), ("end", None)):
            if kind:
                el = self.Ox("w:fldChar")
                el.set(self.qn("w:fldCharType"), kind)
                run._r.append(el)
            elif text is not None:
                if text is instr:
                    el = self.Ox("w:instrText")
                    el.set(self.qn("xml:space"), "preserve")
                    el.text = instr
                    run._r.append(el)
                else:
                    el = self.Ox("w:t")
                    el.text = text
                    run._r.append(el)

    def toc(self):
        para = self.d.add_paragraph()
        self.field(para, 'TOC \\o "1-3" \\h \\z \\u', "Table of contents: in Word, right-click here and choose Update Field.")
        settings = self.d.settings.element
        upd = self.Ox("w:updateFields")
        upd.set(self.qn("w:val"), "true")  # Word offers to fill the table of contents when the file opens
        settings.append(upd)

    def footer(self, text):
        for s in self.d.sections:
            fp = s.footer.paragraphs[0]
            fp.text = ""
            r = fp.add_run(text + "    Page ")
            r.font.size = self.Pt(8)
            r.font.color.rgb = self.RGB.from_string("898781")
            self.field(fp, "PAGE", "1")
            fp.alignment = self.ALIGN.CENTER

    def page_break(self):
        self.d.add_page_break()

    def save(self, path):
        self.d.save(str(path))


# --------------------------------------------------------------------------- the report

def build_report(out: str | Path, author: str | None = None, space_url: str | None = None) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    hr, det, cnn, comp = (_load(n) for n in ("hr_validation.json", "detector.json", "cnn.json", "compression_study.json"))
    tmp = Path(tempfile.mkdtemp(prefix="pulsecheck_doc_"))
    F = Figs(tmp)
    D = Doc()
    today = time.strftime("%d %B %Y")

    # ---------------- title page
    for _ in range(5):
        D.p("")
    D.p("PulseCheck", size=30, color="1C5CAB", align="center").runs[0].bold = True
    D.p("Heart-rate estimation and deepfake screening from the pulse in ordinary face video", size=14, color="3B3A37", align="center")
    D.p("")
    D.p("Technical report: design, implementation and evaluation", size=11.5, color="52514E", align="center")
    D.p("")
    if author:
        D.p(author, size=12, align="center")
    D.p(f"Version {__version__}  ·  generated {today}", size=10, color="898781", align="center")
    if space_url:
        D.link(D.p("Live demo: ", size=10, align="center"), space_url, space_url)
    D.p("")
    status = []
    status.append("heart-rate validation on UBFC-rPPG: " + ("included" if hr else "not run yet"))
    status.append("classical detector on FaceForensics++: " + ("included" if det else "not run yet"))
    status.append("CNN detector on FaceForensics++: " + ("included" if cnn else "not run yet"))
    status.append("compression study: " + ("included" if comp else "not run yet"))
    D.note("Results in this edition:", "; ".join(status) + ". Every result is computed from the real datasets by the commands "
           "listed in Chapter 11; the report is regenerated with `python -m pulsecheck report-doc`.", fill="E9F2FC")
    D.page_break()
    D.h("Contents", 1)
    D.toc()
    D.page_break()

    # ---------------- 1 summary
    D.h("1. Summary", 1)
    D.p("PulseCheck measures a person's heart rate from an ordinary video of their face, without touching them, and uses the "
        "same pulse signal to screen videos for deepfakes. Every heartbeat pushes blood into the skin; blood absorbs green "
        "light, so the skin reflects slightly less green for a moment. A camera records this change. Measuring a pulse with "
        "light is photoplethysmography (PPG); doing it from video at a distance is remote PPG (rPPG).")
    D.p("The deepfake idea follows FakeCatcher (Ciftci, Demir and Yin, 2020): on a real face every patch of skin is fed by the "
        "same heart, so the patches' pulses rise and fall together. A face synthesised or swapped by a generator has no heart "
        "of its own, so its pulse tends to be missing, smeared, or out of step with the skin the generator did not touch.")
    D.p("The project has four parts:")
    cnn_how = (f"trained on the laptop's {cnn['training']['device']} in {cnn['training']['train_minutes']} minutes" if cnn
               else "designed to train on the laptop's NVIDIA RTX 3050 GPU")
    web_how = (f"deployed on Hugging Face Spaces at {space_url}" if space_url
               else "packaged with Docker for deployment on Hugging Face Spaces")
    D.bullets([
        "**A signal-processing pipeline** (Python): face landmarks, skin patches, three classic rPPG algorithms (Green, CHROM, "
        "POS), filtering, Fourier analysis and heart-rate estimation.",
        "**Two detectors, both trained locally**: classical scikit-learn models (CPU) on 91 hand-designed patch-agreement "
        f"features, and a convolutional neural network (PyTorch, {cnn_how}) on FakeCatcher-style PPG maps. Both use the "
        "same splits and are compared head to head.",
        "**An evaluation protocol designed to avoid flattering numbers**: official subject-disjoint splits, grouped "
        "cross-validation, thresholds frozen on validation data, cross-dataset tests and a false-alarm test on genuine videos.",
        f"**A web application** (FastAPI + a static site with interactive Plotly charts), {web_how}, where anyone can "
        "upload or record a clip and see every step of the analysis.",
    ])
    if hr or det or cnn:
        D.h("Headline results", 2)
        rows = []
        if hr:
            hl = hr["headline"]
            rows.append(["Heart rate, UBFC-rPPG", f"MAE {_f(hl['mae'])} bpm, RMSE {_f(hl['rmse'])} bpm, r = {_f(hl['pearson_r'])}",
                         f"{hr['n_videos']} videos, {hl['window_s']} s windows, {METHOD[hl['method']]}"])
        for rep, name in ((det, "Classical detector"), (cnn, "CNN detector")):
            if rep:
                tv = rep["test_video"]
                rows.append([f"{name}, FF++ test", f"AUC {_f(tv.get('auc'), 3)}, balanced accuracy {_pct(tv.get('balanced_accuracy'))}",
                             f"{tv['n']} held-out videos"])
                for g in rep.get("generalization", []):
                    if g["kind"] == "cross_dataset":
                        rows.append([f"{name}, {g['name']}", f"AUC {_f(g.get('auc'), 3)}", f"{g['n_videos']} videos never trained on"])
                    elif g["kind"] == "real_only":
                        rows.append([f"{name}, {g['name']} genuine videos", f"false-alarm rate {_pct(g.get('false_alarm_rate'))}", f"{g['n_videos']} videos"])
        D.table(["Test", "Result", "Data"], rows, widths=[2.0, 2.6, 1.9], caption="Headline results (all on data never used for training).")
    else:
        D.note("Results pending.", "Chapters 7 to 9 are filled in automatically from the real datasets once the commands "
               "in Chapter 11 have been run on them.")

    # ---------------- 2 background
    D.h("2. Background", 1)
    D.h("2.1 Why skin colour carries a pulse", 2)
    D.p("Light entering the skin is partly absorbed by haemoglobin. With each heartbeat the volume of blood in the "
        "microvessels rises and falls, so the amount of light absorbed rises and falls too. Haemoglobin absorbs green light "
        "most strongly, which is why the green channel of a camera carries the clearest pulse (Verkruysse et al., 2008). The "
        "change is tiny: typically well under one percent of the pixel value, i.e. a fraction of one grey level. It is "
        "recoverable only because thousands of pixels are averaged, which cancels sensor noise.")
    D.h("2.2 Why this can expose a deepfake", 2)
    D.p("A deepfake generator is trained to make faces look right frame by frame; it is not trained to reproduce the "
        "sub-grey-level, spatially coherent colour rhythm of blood flow. FakeCatcher used this: it extracted pulse signals from "
        "32 face regions, arranged them into 'PPG maps', and trained a CNN to tell real from fake, reporting about 96% "
        "accuracy on an older benchmark and about 91% on its own harder set.")
    D.note("Important caveat.", "Seibold et al. (2025, Frontiers in Imaging) showed that high-quality face swaps can inherit "
           "the pulse of the source video they were made from. A believable pulse is therefore evidence, not proof, and "
           "this project presents its output as a screening signal.")

    # ---------------- 3 system
    D.h("3. System overview", 1)
    D.figure(F.pipeline(), "The pipeline. The top row is shared; the classical and the CNN detector branch off the colour traces.")
    D.table(["Stage", "Module", "What it does"], [
        ["1 Read video", "pulsecheck/video.py", "Decodes frames with OpenCV and recovers a trustworthy timestamp for each frame."],
        ["2 Find the face", "pulsecheck/face.py", "MediaPipe Face Landmarker places 478 3-D points on the face in every frame."],
        ["3 Skin patches", "pulsecheck/roi.py", "Six named regions plus a 32-patch grid; a per-video skin-colour model removes hair and brows."],
        ["4 Traces", "pulsecheck/pipeline.py, traces.py", "Mean R, G, B of every patch per frame, resampled to a uniform 30 Hz clock and cached."],
        ["5 Pulse", "pulsecheck/rppg.py", "Green, CHROM and POS turn three colour traces into one pulse waveform."],
        ["6 Heart rate", "pulsecheck/signals.py", "Detrend, band-pass 0.7-3 Hz, Fourier transform, peak, SNR."],
        ["7a Features", "pulsecheck/features.py", "91 numbers per 8 s window describing how well the patches agree."],
        ["7b PPG maps", "pulsecheck/deep.py", "2 x 32 x 224 images of the grid patches' pulses for the CNN."],
        ["8 Decide", "detector.py, deep.py", "Each window gets P(fake); the video score is their mean; a threshold gives the verdict."],
        ["9 Evaluate", "hr_eval.py, detect_eval.py, studies.py", "Protocols that write reports/*.json for the website and this report."],
    ], widths=[1.1, 1.9, 3.5], caption="Stages and the code that implements them.")

    # ---------------- 4 implementation
    D.h("4. Implementation", 1)
    D.h("4.1 Reading video and keeping time", 2)
    D.p("Every beats-a-minute number depends on knowing when each frame was captured. Files from phones and from browsers "
        "(WebM recorded with MediaRecorder) are often variable-frame-rate and report a meaningless nominal frame rate, so "
        "the reader collects both the decoder's presentation timestamp and frame-index/FPS for every frame, uses the "
        "presentation timestamps when they are complete and strictly increasing, and otherwise falls back to the index "
        "clock. All traces are then resampled by linear interpolation onto a uniform 30 Hz grid, so the rest of the code "
        f"never deals with irregular timing. Frames larger than 720 px on the long side are shrunk with area interpolation "
        "first; averaging a patch is almost unaffected by this, and landmarking is faster.")
    D.h("4.2 Face landmarks", 2)
    D.p("MediaPipe Face Landmarker (Tasks API, model face_landmarker.task, float16, SHA-256 verified on download) returns 468 "
        "face-mesh points plus 10 iris points. It runs in VIDEO mode, which tracks the face between frames instead of "
        "re-detecting it every frame (faster and steadier). Frames without a face are marked missing and filled by "
        "interpolation; a clip is rejected if a face is found in fewer than 60% of frames.")
    D.h("4.3 Skin patches and the skin-colour model", 2)
    D.p("Two sets of patches are measured in every frame:")
    D.bullets([
        "**Six named regions** for heart rate and the hand-made features: left and right forehead, left and right cheek, nose "
        "bridge and chin. Each is a polygon through fixed landmark numbers, chosen by plotting all landmark numbers on a face "
        "(`python -m pulsecheck landmarks`) and shrunk 8-12% towards its centre so its edges stay off brows and hairline.",
        f"**A 32-patch grid** for the CNN: small squares (side = {0.14:.2f} x the outer-eye-corner distance) centred on 32 "
        "landmarks across forehead, both cheeks, nose bridge and chin, ordered so neighbouring rows are neighbouring skin.",
    ])
    D.table(["Region", "Landmark polygon"], [[n.replace("_", " "), ", ".join(map(str, ROI_LANDMARKS[n]))] for n in ROI_NAMES],
            widths=[1.4, 5.1], caption="The six named regions (left/right are the subject's own).")
    D.p("Inside each patch, hair, eyebrows, beard and glasses frames must not be averaged with skin. Fixed skin-colour "
        "thresholds are known to work worse on darker skin, so instead PulseCheck learns a skin model from each video: it "
        "fits a 2-D Gaussian to the chrominance (Cr, Cb in YCrCb space, ignoring brightness) of the cheek pixels in the first "
        "ten frames, after trimming the 5% most extreme samples, and then keeps a pixel only if its squared Mahalanobis "
        "distance is below 13.8 (the 99.9% quantile of a chi-squared distribution with 2 degrees of freedom). If a mask would "
        "leave less than a quarter of a patch, the whole patch is used rather than failing.")
    D.h("4.4 From colour to pulse: Green, CHROM and POS", 2)
    D.p("Each patch gives three numbers per frame: mean red, green and blue, written C(t) = [R, G, B]. Lighting changes and "
        "motion change all three channels together and are usually far larger than the pulse; the algorithms differ in how "
        "they cancel them. All outputs are oriented so that the waveform goes up when blood volume goes up.")
    D.table(["Method", "Idea", "Computation"], [
        ["Green (Verkruysse 2008)", "Use the channel blood affects most.", "p(t) = -(G(t)/mean(G) - 1)"],
        ["CHROM (de Haan & Jeanne 2013)", "Two colour-difference signals in which brightness cancels, mixed with a tuned weight.",
         "Cn = C/mean(C); X = 3Rn - 2Gn; Y = 1.5Rn + Gn - 1.5Bn; band-pass both; S = Xf - (sd(Xf)/sd(Yf)) * Yf; 1.6 s Hann windows, overlap-added"],
        ["POS (Wang et al. 2017)", "Project colour onto a plane orthogonal to the skin tone, then combine the two axes.",
         "Cn = C/mean(C) per 1.6 s window; S1 = Gn - Bn; S2 = -2Rn + Gn + Bn; h = S1 + (sd(S1)/sd(S2)) * S2; overlap-add"],
    ], widths=[1.5, 2.0, 3.0], caption="The three rPPG algorithms implemented in pulsecheck/rppg.py.")
    D.p(f"The {RPPG_WINDOW_S} s window is about one heartbeat long, so the weights can follow slow changes in lighting. The unit "
        "tests include a case where a lamp flickers at 1.8 Hz with twice the strength of the pulse: POS and CHROM still find "
        "the true rate while Green locks onto the flicker, which is exactly the reason those methods exist.")
    D.h("4.5 Filtering and reading the beats", 2)
    D.bullets([
        "**Detrend**: subtract the best straight line (scipy.signal.detrend) to remove slow drift.",
        f"**Band-pass**: zero-phase 3rd-order Butterworth filter (second-order sections, sosfiltfilt) keeping {HR_BAND[0]}-{HR_BAND[1]} Hz, "
        f"i.e. {HR_BAND[0] * 60:.0f}-{HR_BAND[1] * 60:.0f} beats a minute. Zero-phase means the pulse is not shifted in time.",
        "**Spectrum**: Hann-windowed periodogram, zero-padded 8x. A clip of length T only resolves frequencies 1/T apart "
        "(a 10 s clip at 30 fps is 300 numbers with steps 0.1 Hz = 6 bpm apart; 72 bpm = 1.2 Hz sits on step 12). Zero-padding "
        "adds no information, but interpolates the spectrum so a single dominant peak can be located between steps; the "
        "peak is refined further with a three-point parabola.",
        "**Heart rate** = 60 x the strongest frequency inside the band.",
        "**Signal quality (SNR)** after de Haan & Jeanne (2013): power within +/-0.1 Hz of the heart rate and +/-0.2 Hz of its "
        "first harmonic, divided by all other power between 0.5 and 4 Hz, in decibels. 0 dB means as much power off the "
        "pulse as on it.",
    ])
    D.h("4.6 Patch-agreement features (classical detector)", 2)
    names = feature_names(ROI_NAMES)
    D.p(f"Each video is cut into {DETECT_WINDOW_S:.0f} s windows with a {DETECT_STEP_S:.0f} s hop. Pulses are computed once "
        "per patch over the whole clip (so filters have no edge effects inside a window) and then sliced. Each window becomes "
        f"a row of {len(names)} numbers:")
    D.table(["Family", "Count", "Definition"], [
        ["Global", "10", "pooled SNR and spectral entropy; spread and range of the six per-patch heart rates; mean and minimum "
         "pairwise coherence and correlation; mean absolute lag; log ratio of 4-12 Hz to heart-band power (FakeCatcher's wide band)"],
        ["Per patch (x6)", "36", "SNR at the pooled heart rate; gap between the patch's own rate and the pooled rate; correlation "
         "with the mean of the other five patches; high-band ratio; log amplitude relative to the median patch; log std of 4-12 Hz green flicker"],
        ["Per pair (x15)", "45", "magnitude-squared coherence in the heart band, weighted by sqrt(Pxx*Pyy); Pearson correlation; "
         "lag of the best cross-correlation within +/-0.33 s"],
    ], widths=[1.2, 0.6, 4.7], caption="Feature families (full list in Appendix A).")
    D.p("Coherence deserves a note. Averaging coherence evenly across the 0.7-3 Hz band is dominated by the many frequencies "
        "that contain only noise and gave values around 0.3 even for patches that visibly share a beat. Weighting each "
        "frequency by the geometric mean of the two patches' own power concentrates the average where both patches carry "
        "energy, which on a real face is the heart rate, without favouring frequencies merely because the patches agree there.")
    D.h("4.7 Classical classifiers", 2)
    D.table(["Model", "Configuration"], [
        ["Logistic regression", "median imputation, standardisation, L2, C = 0.3, class_weight = balanced"],
        ["SVM (RBF)", "median imputation, standardisation, C = 1, gamma = scale, Platt-scaled probabilities, balanced"],
        ["Gradient boosting", "HistGradientBoosting, 300 iterations, learning rate 0.05, 15 leaves, L2 = 1, balanced"],
        ["Random forest", "400 trees, min 3 samples per leaf, balanced_subsample class weights"],
    ], widths=[1.6, 4.9], caption="Candidate models (scikit-learn); the best by grouped cross-validation is kept.")
    D.h("4.8 CNN on PPG maps (trained on the GPU)", 2)
    D.p(f"For each window the 32 grid patches' POS pulses form an image: 32 rows (patches) by {int(DETECT_WINDOW_S * TARGET_FS)} "
        "columns (samples). Channel 0 is the heart band (0.7-3 Hz) and channel 1 the wide band "
        f"({WIDE_BAND[0]}-{WIDE_BAND[1]:.0f} Hz); every row is standardised. Training uses a random 224-sample crop of each "
        "window (a time-shift augmentation) and random sign flips (the polarity of a pulse says nothing about fakeness); "
        "evaluation uses the centre crop.")
    D.table(["Layer", "Output shape", "Parameters"], [
        ["Input", "2 x 32 x 224", "-"],
        ["Conv 3x9 -> BatchNorm -> ReLU -> MaxPool 1x2", "16 x 32 x 112", "16 x 2 x 27 + 32"],
        ["Conv 3x7 -> BatchNorm -> ReLU -> MaxPool 2x2", "32 x 16 x 56", "32 x 16 x 21 + 64"],
        ["Conv 3x5 -> BatchNorm -> ReLU -> MaxPool 2x2", "64 x 8 x 28", "64 x 32 x 15 + 128"],
        ["Conv 3x3 -> BatchNorm -> ReLU", "64 x 8 x 28", "64 x 64 x 9 + 128"],
        ["Global average pool -> Dropout 0.3 -> Linear", "1 logit", "65"],
    ], widths=[3.4, 1.4, 1.7], caption="CNN architecture (about 80,000 parameters).")
    D.p("Training: AdamW (learning rate 1e-3, weight decay 1e-2) with a cosine schedule, batch 256, binary cross-entropy "
        "with a positive-class weight equal to the real/fake ratio (FaceForensics++ has five fakes per real video), automatic "
        "mixed precision (float16) on CUDA, and early stopping when the validation video-level AUC has not improved for 8 "
        "epochs (maximum 40). The weights with the best validation AUC are kept. A model this size fits easily in the "
        "RTX 3050's 4 GB; the training windows stay in system memory as float16 and are moved to the GPU batch by batch.")
    D.h("4.9 From windows to a verdict", 2)
    D.p("Each window's probability of being fake is averaged over the video (every window votes). The decision threshold "
        "is the one that maximises balanced accuracy on the validation split and is then frozen; the website says "
        "'uncertain' when the score is within 0.1 of it. The website serves whichever detector, classical or CNN, has the "
        "higher validation AUC; the test set is never used to choose.")

    # ---------------- 5 datasets
    D.h("5. Datasets", 1)
    D.table(["Dataset", "Contents", "Use here", "Access and terms"], [
        ["UBFC-rPPG, DATASET_2 (Bobbia et al., 2019)", "42 uncompressed 640x480 30 fps webcam videos of people doing a maths game, "
         "with a finger pulse-oximeter PPG recorded alongside", "heart-rate error; false alarms on genuine videos",
         "public Drive folder; Agreement.xlsx; facial images of subjects 21 and 27 must not be published"],
        ["FaceForensics++ c23 (Rossler et al., 2019)", "1,000 YouTube videos and 5 manipulations of each: Deepfakes, Face2Face, "
         "FaceSwap, NeuralTextures, FaceShifter", "train / validate / test the detectors", "request form; terms of use; no redistribution"],
        ["Celeb-DF v2 (Li et al., 2020)", "590 celebrity + 300 YouTube real videos, 5,639 high-quality fakes; official 518-video test list",
         "cross-dataset test only", "request form; terms of use; no redistribution"],
    ], widths=[1.5, 2.1, 1.3, 1.6], caption="Datasets. None is included in the repository or the deployed site.", font=8.5)
    D.p("FaceForensics++ ships official train/validation/test lists of source-target pairs (720/140/140 videos). Both "
        "videos of a pair and all their fakes share a group, so no person appears on both sides of a split. Only the first "
        "20 s of each FaceForensics++ and Celeb-DF video are processed; UBFC videos are processed whole.")

    # ---------------- 6 protocol
    D.h("6. Evaluation protocol", 1)
    D.bullets([
        "**No person on both sides.** Model selection uses 5-fold GroupKFold inside the training split; groups are FF++ source "
        "pairs. A random split of windows would put windows of the same video in training and testing; the report measures "
        "how much that inflates the score.",
        "**Decisions frozen before testing.** The model, the early-stopping epoch and the threshold are all chosen on "
        "training or validation data only.",
        "**Video-level scores** are the headline; window-level scores are reported alongside.",
        "**Generalisation.** The FF++-trained models are tested unchanged on Celeb-DF (a different dataset with "
        "higher-quality fakes) and on UBFC-rPPG's genuine webcam videos, where any 'fake' verdict is a false alarm.",
        "**Heart-rate ground truth** for a window comes from the oximeter's PPG waveform analysed with exactly the same "
        "spectral method, as the dataset's authors recommend, so both sides have the same frequency resolution.",
    ])
    D.table(["Metric", "Meaning"], [
        ["MAE / RMSE (bpm)", "mean absolute / root-mean-square difference between video and oximeter heart rate"],
        ["Pearson r", "linear agreement of the two heart-rate series"],
        ["Bland-Altman bias, limits", "mean difference, and the range holding 95% of differences (bias +/- 1.96 SD)"],
        ["Within 5 bpm", "share of windows whose estimate is within 5 bpm of the oximeter"],
        ["ROC AUC", "probability a random fake scores above a random real video; 0.5 is chance, 1.0 perfect (95% CI by bootstrap)"],
        ["Balanced accuracy", "mean of the detection rate on fakes and on reals; robust to the 5:1 class imbalance"],
        ["EER", "error rate at the threshold where false alarms equal misses"],
        ["Average precision", "area under the precision-recall curve"],
        ["False-alarm rate", "share of genuine videos called fake"],
    ], widths=[1.8, 4.7], caption="Metrics.")

    # ---------------- 7 results: heart rate
    D.h("7. Results: heart rate", 1)
    if hr:
        hl = hr["headline"]
        D.p(f"On {hr['n_videos']} UBFC-rPPG videos, {METHOD[hl['method']]} with {hl['window_s']} s windows gives a mean absolute "
            f"error of **{_f(hl['mae'])} bpm** (RMSE {_f(hl['rmse'])} bpm, Pearson r = {_f(hl['pearson_r'])}); "
            f"{_pct(hl['within_5bpm'], 0)} of windows are within 5 bpm of the oximeter. Bland-Altman bias is {_f(hl['bias'])} bpm "
            f"with 95% limits of agreement {_f(hl['loa_low'], 1)} to {_f(hl['loa_high'], 1)} bpm.")
        D.figure(F.hr_scatter_ba(hr), f"Video estimate against the finger oximeter, {METHOD[hl['method']]}, {hl['window_s']} s windows.")
        rows = [[METHOD[s["method"]], s["window_s"], s["n"], _f(s["mae"]), _f(s["rmse"]), _f(s["pearson_r"]), _pct(s["within_5bpm"], 0),
                 _f(s.get("snr_mean"), 1)] for s in hr["summary"]]
        D.table(["Method", "Window (s)", "Windows", "MAE", "RMSE", "r", "<=5 bpm", "SNR (dB)"], rows, caption="Heart-rate error by method and window length.")
        D.figure(F.hr_windows(hr), "Error against window length. Longer windows resolve frequency more finely.", width=4.8)
        D.figure(F.hr_subjects(hr), "Mean absolute error per person, sorted.")
        skin = [b for b in hr["by_skin"] if b["n_subjects"]]
        D.table(["ITA category", "People", "MAE (bpm)"], [[b["category"], b["n_subjects"], _f(b["mae"])] for b in hr["by_skin"]],
                caption="Error by estimated skin-tone angle (ITA). ITA from video is approximate and lighting-dependent.")
        if len(skin) <= 2:
            D.note("Fairness caveat.", "UBFC-rPPG covers too narrow a range of skin tones to draw conclusions about darker skin; "
                   "that gap in the data is itself the finding.")
    else:
        D.pending("Heart-rate error against the pulse oximeter", [
            "python -m pulsecheck extract ubfc --root C:\\datasets\\UBFC-rPPG --max-seconds 0",
            "python -m pulsecheck eval-hr"])
    D.h("7.1 Compression", 2)
    if comp:
        D.figure(F.compression(comp), "Heart-rate error when the same videos are re-encoded with H.264 at increasing CRF.", width=4.8)
        D.table(["Encoding", "kbit/s", "POS MAE 10 s", "POS MAE 30 s", "Green MAE 10 s"],
                [[r["label"], _f(r["kbps"], 0), _f(r["mae_pos_10s"]), _f(r["mae_pos_30s"]), _f(r["mae_green_10s"])] for r in comp["rows"]],
                caption=f"Compression study on {comp['n_videos']} UBFC-rPPG videos (first {comp['max_s']:.0f} s).")
    else:
        D.pending("Heart-rate error at several H.264 compression levels", [
            "python -m pulsecheck compression-study --root C:\\datasets\\UBFC-rPPG --limit 20"])

    # ---------------- 8 results: detectors
    D.h("8. Results: deepfake detection", 1)
    if det or cnn:
        if det:
            D.table(["Model", "Window AUC (CV)", "Video AUC (CV)", "Seconds"],
                    [[m["label"], _f(m["window_auc"], 3), _f(m["video_auc"], 3), m["seconds"]] for m in det["model_selection"]],
                    caption="Classical model selection: 5-fold GroupKFold on the training split.")
        rows = []
        for rep, name in ((det, det and det.get("chosen_label")), (cnn, "CNN on PPG maps")):
            if rep:
                tv, tw = rep["test_video"], rep["test_window"]
                ci = tv.get("auc_ci95") or [None, None]
                rows.append([name, _f(rep.get("val_video_auc"), 3), f"{_f(tv.get('auc'), 3)} [{_f(ci[0], 2)}-{_f(ci[1], 2)}]",
                             _pct(tv.get("balanced_accuracy")), _pct(tv.get("eer")), _f(tv.get("average_precision"), 3), _f(tw.get("auc"), 3)])
        D.table(["Model", "Val AUC", "Test AUC [95% CI]", "Bal. acc.", "EER", "AP", "Window AUC"], rows,
                caption="Held-out FaceForensics++ test videos (video level unless stated).")
        D.figure(F.rocs(det, cnn), "Left: test ROC curves. Right: real test videos against each manipulation method.")
        src = det or cnn
        D.table(["Manipulation", "AUC", "Fakes caught at threshold", "Fake videos"],
                [[p["method"], _f(p["auc"], 3), _pct(p["detection_rate"]), p["n_fake"]] for p in src["per_method"]],
                caption="Per manipulation method.")
        if det:
            c = det["test_video"]["confusion"]
            D.table(["", "Called real", "Called fake"], [["Actually real", c["tn"], c["fp"]], ["Actually fake", c["fn"], c["tp"]]],
                    widths=[1.6, 1.4, 1.4], caption=f"Confusion matrix, classical detector, threshold {_f(det['threshold'], 3)} chosen on validation.")
            sd = det["split_demo"]
            D.table(["Split", "Window AUC", "Video AUC"], [
                ["Grouped by person (honest)", _f(sd["grouped_split"]["window_auc"], 3), _f(sd["grouped_split"]["video_auc"], 3)],
                ["Random windows (leaky)", _f(sd["random_window_split"]["window_auc"], 3), _f(sd["random_window_split"]["video_auc"], 3)],
            ], caption="The same model scored with an honest and with a leaky split, inside the training data.")
            D.figure(F.importance(det), "Which features the classical model relies on.", width=5.6)
        if cnn:
            tr = cnn["training"]
            D.p(f"The CNN ({tr['parameters']:,} parameters) trained on {tr['device']} for {tr['epochs_run']} epochs "
                f"({tr['train_minutes']} min); the best validation AUC came at epoch {tr['best_epoch']}.")
            D.figure(F.learning(cnn), "CNN training: loss and validation AUC per epoch; the vertical line marks the kept weights.")
    else:
        D.pending("Detector training and testing on FaceForensics++", [
            "python -m pulsecheck extract ffpp --root C:\\datasets\\FaceForensics --compression c23",
            "python -m pulsecheck extract celebdf --root C:\\datasets\\Celeb-DF-v2",
            "python -m pulsecheck train --cross Celeb-DF=cache/celebdf --real UBFC=cache/ubfc",
            "python -m pulsecheck train-cnn --cross Celeb-DF=cache/celebdf --real UBFC=cache/ubfc"])

    # ---------------- 9 generalisation
    D.h("9. Results: generalisation", 1)
    rows = []
    for rep, name in ((det, "classical"), (cnn, "CNN")):
        for g in (rep or {}).get("generalization", []):
            rows.append([name, g["name"], g["kind"].replace("_", " "), g["n_videos"], _f(g.get("auc"), 3), _pct(g.get("false_alarm_rate"))])
    if rows:
        D.table(["Model", "Test set", "Type", "Videos", "AUC", "False-alarm rate"], rows, caption="Data the detectors never saw.")
        D.p("A drop from the FaceForensics++ test set to Celeb-DF is expected: its fakes are of higher quality and its videos "
            "come from a different source. This cross-dataset number, not the home-dataset number, is the honest one to quote.")
    else:
        D.pending("Cross-dataset (Celeb-DF) and genuine-video (UBFC) results", ["see Chapter 8 commands (--cross and --real)"])

    # ---------------- 10 web app + deployment
    D.h("10. Web application and deployment", 1)
    D.h("10.1 Backend", 2)
    D.p("FastAPI (Python) serves the API and the static site from one process. Analysis runs on a single background "
        "thread, because a free Hugging Face CPU Space has 2 virtual CPUs; the browser polls for progress, which is reported "
        "frame by frame during landmarking. Uploads are streamed to a temporary file with a size cap (60 MB), analysed for "
        "at most 20 s, and deleted as soon as the job finishes; nothing is stored or logged.")
    D.table(["Endpoint", "Purpose"], [
        ["POST /api/analyze", "upload a video (multipart) and a method; returns a job id"],
        ["GET /api/jobs/{id}", "status, stage, progress, queue position, and the result when done"],
        ["GET /api/info", "version, loaded models and their validation/test AUC, limits"],
        ["GET /api/reports", "the evaluation reports (real datasets only)"],
        ["GET /healthz", "liveness probe used by Docker"],
        ["GET /api/docs", "interactive OpenAPI documentation"],
    ], widths=[1.8, 4.7], caption="HTTP API.")
    D.h("10.2 Frontend", 2)
    D.p("A single static page in plain HTML, CSS and ES-module JavaScript, with Plotly.js 3.7.0 (from the jsDelivr CDN) for "
        "interactive charts. Four views: Analyze (upload, 15 s webcam recording through MediaRecorder at 6 Mbit/s because "
        "compression destroys the pulse, every step of the analysis as charts, the patch overlay, the verdict), Results (the "
        "evaluation dashboards), How it works (a live in-browser lab for the Fourier maths) and About (model card).")
    D.p("Charts follow one set of rules: a colour palette validated for colour-vision deficiency in light and dark mode, "
        "2 px lines, rounded bar ends, hairline grids, a legend whenever there are two or more series, and a 'Table' button "
        "on every chart that shows the same numbers as an accessible table. Each skin patch keeps the same colour in the "
        "overlay image and in every chart.")
    D.h("10.3 Deployment", 2)
    D.p("A Dockerfile builds on python:3.14-slim, adds the shared libraries OpenCV and MediaPipe need (libgl1, libglib2.0-0, "
        "libportaudio2), installs the pinned requirements as a non-root user with uid 1000 (as Hugging Face requires), "
        "downloads and checksum-verifies the face model, and starts Uvicorn on port 7860 with a health check. The README "
        "front matter (sdk: docker, app_port: 7860) configures the Hugging Face Space, and the Hugging Face CLI uploads "
        "the project (it stores binary files with Xet automatically). Training always happens on the laptop; only the "
        "trained model files are uploaded.")

    # ---------------- 11 reproduce
    D.h("11. How to reproduce", 1)
    D.code("""\
python -m venv .venv  &&  .venv\\Scripts\\activate
pip install -r requirements-dev.txt
pip install torch==2.14.1 --index-url https://download.pytorch.org/whl/cu126     # GPU training
python -m pulsecheck download-model
python -m pytest                                                            # all tests

python -m pulsecheck extract ubfc   --root C:\\datasets\\UBFC-rPPG --max-seconds 0
python -m pulsecheck eval-hr
python -m pulsecheck compression-study --root C:\\datasets\\UBFC-rPPG --limit 20
python -m pulsecheck extract ffpp   --root C:\\datasets\\FaceForensics --compression c23
python -m pulsecheck extract celebdf --root C:\\datasets\\Celeb-DF-v2
python -m pulsecheck train      --cross Celeb-DF=cache/celebdf --real UBFC=cache/ubfc
python -m pulsecheck train-cnn  --cross Celeb-DF=cache/celebdf --real UBFC=cache/ubfc
python -m pulsecheck report-doc --author "Your Name"
python -m pulsecheck serve                                                  # http://127.0.0.1:7860""")

    # ---------------- 12 testing
    D.h("12. Testing", 1)
    tests = sorted((ROOT / "tests").glob("test_*.py"))
    n_tests = sum(len(re.findall(r"^def test_", t.read_text(encoding="utf-8"), re.M)) for t in tests)
    D.p(f"The repository has {n_tests} automated test functions (pytest) in {len(tests)} files. They check the code against answers "
        "known in advance: the article's worked example (300 samples, step 12, 72 bpm), heart-rate recovery at off-bin "
        "rates, filter behaviour, all three rPPG methods (including the lamp-flicker case), coherence and features, the "
        "dataset loaders on miniature folder trees laid out like the official downloads, the metrics against hand-computed "
        "values, the web API end to end, and the CNN on a learnable toy task.")
    D.note("About synthetic data.", "To test the pipeline end to end, the test suite renders short clips with a known pulse "
           "painted onto a public-domain photograph. These exist only to check that the code computes the right answer; "
           "they are never used for any reported result, model, or anything shown on the website.")

    # ---------------- 13 limitations
    D.h("13. Limitations and responsible use", 1)
    D.bullets([
        "Not a forensic tool. A verdict is a screening signal and must never be used to accuse anyone.",
        "High-quality face swaps can inherit their source video's pulse (Seibold et al., 2025).",
        "Compression, head movement, uneven or flickering light, low resolution and heavy make-up all weaken the signal.",
        "Darker skin absorbs more light, which lowers the pulse's strength; UBFC-rPPG has few darker-skinned subjects, so "
        "accuracy there is not established by this work.",
        "The detectors are trained on YouTube-style FaceForensics++ videos; webcam clips and new generators are different "
        "distributions, so performance on them is less certain than the test numbers.",
        "Privacy: the website processes uploads in memory and deletes them; the datasets are never redistributed, and the "
        "UBFC restriction on subjects 21 and 27 is respected.",
    ])

    # ---------------- 14 stack
    D.h("14. Technology stack", 1)
    stack = [
        ("Python", platform.python_version(), "language for the pipeline, training and server"),
        ("NumPy", _ver("numpy"), "arrays and vectorised signal maths"),
        ("SciPy", _ver("scipy"), "Butterworth filters, periodograms, Welch coherence, detrending, statistics"),
        ("OpenCV", _ver("opencv-contrib-python"), "video decoding (FFmpeg backend), colour spaces, drawing"),
        ("MediaPipe", _ver("mediapipe"), "Face Landmarker: 478 facial landmarks per frame"),
        ("scikit-learn", _ver("scikit-learn"), "classical classifiers, GroupKFold, metrics, permutation importance"),
        ("PyTorch", _ver("torch"), "the CNN, trained with CUDA on an NVIDIA RTX 3050 (4 GB)"),
        ("joblib", _ver("joblib"), "saving the classical model"),
        ("FastAPI", _ver("fastapi"), "web API"),
        ("Uvicorn", _ver("uvicorn"), "ASGI web server"),
        ("python-multipart", _ver("python-multipart"), "video uploads"),
        ("Plotly.js", "3.7.0", "interactive charts in the browser"),
        ("pytest", _ver("pytest"), "tests"),
        ("imageio-ffmpeg", _ver("imageio-ffmpeg"), "H.264 encoding for the compression study and test clips"),
        ("matplotlib", _ver("matplotlib"), "figures in this report"),
        ("python-docx", _ver("python-docx"), "this report"),
        ("Docker", "python:3.14-slim", "container image"),
        ("Hugging Face Spaces", "Docker SDK", "public hosting"),
    ]
    D.table(["Component", "Version", "Used for"], [[a, b or "not installed", c] for a, b, c in stack],
            widths=[1.5, 1.2, 3.8], caption="Software used (versions read from the installed packages).")

    # ---------------- 15 glossary + references
    D.h("15. Glossary", 1)
    D.table(["Term", "Meaning"], [
        ["PPG / rPPG", "photoplethysmography: measuring the pulse with light; remote PPG does it from video"],
        ["ROI / patch", "region of interest: an area of skin whose pixels are averaged"],
        ["Band-pass filter", "keeps only a range of frequencies (here where a heartbeat can be)"],
        ["Fourier transform / FFT", "splits a signal into the frequencies inside it"],
        ["SNR", "signal-to-noise ratio: power on the pulse versus power elsewhere"],
        ["Coherence", "how consistently two signals move together at a given frequency (0 to 1)"],
        ["PPG map", "an image whose rows are patches and columns are time, holding each patch's pulse"],
        ["CNN", "convolutional neural network"],
        ["GroupKFold", "cross-validation that keeps every group (person) within one fold"],
        ["AUC", "area under the ROC curve"],
        ["ITA", "individual typology angle: a skin-tone measure from CIELAB colour"],
    ], widths=[1.7, 4.8], caption="Terms.")
    D.h("16. References", 1)
    refs = [
        ("W. Verkruysse, L. O. Svaasand, J. S. Nelson. Remote plethysmographic imaging using ambient light. Optics Express 16(26), 2008.", "https://doi.org/10.1364/OE.16.021434"),
        ("G. de Haan, V. Jeanne. Robust pulse rate from chrominance-based rPPG. IEEE Trans. Biomedical Engineering 60(10), 2013.", "https://doi.org/10.1109/TBME.2013.2266196"),
        ("W. Wang, A. C. den Brinker, S. Stuijk, G. de Haan. Algorithmic principles of remote PPG. IEEE Trans. Biomedical Engineering 64(7), 2017.", "https://doi.org/10.1109/TBME.2016.2609282"),
        ("U. A. Ciftci, I. Demir, L. Yin. FakeCatcher: Detection of synthetic portrait videos using biological signals. IEEE TPAMI, 2020.", "https://arxiv.org/abs/1901.02212"),
        ("S. Bobbia, R. Macwan, Y. Benezeth, A. Mansouri, J. Dubois. Unsupervised skin tissue segmentation for remote photoplethysmography. Pattern Recognition Letters 124, 2019.", "https://doi.org/10.1016/j.patrec.2017.10.017"),
        ("A. Rossler et al. FaceForensics++: Learning to detect manipulated facial images. ICCV, 2019.", "https://arxiv.org/abs/1901.08971"),
        ("Y. Li, X. Yang, P. Sun, H. Qi, S. Lyu. Celeb-DF: A large-scale challenging dataset for DeepFake forensics. CVPR, 2020.", "https://arxiv.org/abs/1909.12962"),
        ("C. Seibold, E. L. Wisotzky, A. Beckmann, B. Kossack, A. Hilsmann, P. Eisert. High-quality deepfakes have a heart! Frontiers in Imaging 4, 2025.", "https://doi.org/10.3389/fimag.2025.1504551"),
        ("X. Liu et al. rPPG-Toolbox: Deep remote PPG toolbox. NeurIPS Datasets and Benchmarks, 2023.", "https://arxiv.org/abs/2210.00716"),
    ]
    for i, (text, url) in enumerate(refs, 1):
        para = D.p(f"[{i}] {text} ")
        D.link(para, url, url)

    # ---------------- appendices
    D.page_break()
    D.h("Appendix A. All detector features", 1)
    D.table(["#", "Name", "Description"], [[i + 1, n, describe(n)] for i, n in enumerate(names)], widths=[0.4, 3.6, 2.5],
            caption=f"The {len(names)} features of one window.", font=8)
    D.h("Appendix B. The 32-patch grid", 1)
    D.p("Landmark numbers in row order: " + ", ".join(map(str, GRID_LANDMARKS)) + ".")
    D.h("Appendix C. Command reference", 1)
    from .cli import build_parser

    sub = next(a for a in build_parser()._actions if a.__class__.__name__ == "_SubParsersAction")
    D.table(["Command", "What it does"], [[c.dest, c.help] for c in sub._choices_actions], widths=[1.6, 4.9],
            caption="python -m pulsecheck <command> --help shows every option.")
    models = [json.loads(p.read_text()) for p in (MODELS_DIR / "detector.json", MODELS_DIR / "cnn.json") if p.exists()]
    if models:
        D.h("Appendix D. Trained model files", 1)
        D.table(["Model", "Trained", "On", "Val AUC", "Test AUC", "Threshold"],
                [[m.get("model_label"), m.get("trained_at"), m.get("train_source"), _f(m.get("val_video_auc"), 3),
                  _f(m.get("test_video_auc"), 3), _f(m.get("threshold"), 3)] for m in models], caption="models/*.json")

    D.footer(f"PulseCheck technical report · v{__version__}")
    D.save(out)
    return out
